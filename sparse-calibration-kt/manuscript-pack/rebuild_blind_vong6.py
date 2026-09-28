#!/usr/bin/env python3
"""Rebuild the blind manuscript from the Word-normalized named file."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from apply_vong6_layout import export_one  # noqa: E402
from manuscript_paths import BLIND_DOC, BLIND_DOCX, BLIND_PDF, FULL_DOCX  # noqa: E402

# Opening a path ending in _blind.docx can hang Word after a killed session.
REVIEW_DOCX = BLIND_DOCX.with_name("_layout_review.docx")

BLIND_DATA = (
    "Anonymized code, data-preprocessing scripts, frozen diagnostic summaries, and the "
    "accompanying-artifact files needed to rebuild the reported tables (including the "
    "τ-grid and controlled-sparsification grid) are provided as "
    "code_for_review_anonymous.zip for double-blind review. Source benchmark datasets "
    "are publicly released by their original providers. The public repository URL will "
    "be released upon acceptance."
)
CONTRIB = (
    "Author 1: Conceptualization, Methodology, Software, Formal analysis, and Writing "
    "(led diagnostic design, experiments, and manuscript). Authors 2 and 3: Software "
    "(data processing and baseline runs). Author 4: Methodology (methodological review). "
    "Author 5: Supervision and Writing (manuscript revision). All authors approved the "
    "final version."
)


def set_paragraph(p, text: str) -> None:
    runs = list(p.runs)
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def find_one(paras, pred, label: str):
    hits = [i for i, p in enumerate(paras) if pred(p.text or "")]
    if len(hits) != 1:
        raise SystemExit(f"{label} hits {hits}")
    return hits[0]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    shutil.copy2(FULL_DOCX, BLIND_DOCX)
    d = Document(str(BLIND_DOCX))
    paras = d.paragraphs
    author_i = find_one(paras, lambda t: t.startswith("Khanh-Trinh"), "authors")
    set_paragraph(paras[author_i], "Anonymous Authors")
    set_paragraph(paras[author_i + 1], "Affiliations omitted for double-blind review.")
    for idx in range(author_i + 2, author_i + 5):
        set_paragraph(paras[idx], "")
    contrib = paras[find_one(paras, lambda t: t.startswith("K.-T.N.:"), "contributions")]
    set_paragraph(contrib, CONTRIB)
    data = paras[find_one(paras, lambda t: "github.com/trinhnkt" in t, "data")]
    set_paragraph(data, BLIND_DATA)
    ack = paras[find_one(paras, lambda t: "no specific funding" in t, "acknowledgment")]
    set_paragraph(ack, "Omitted for double-blind review.")
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.core_properties.title = d.core_properties.title or ""
    d.save(str(BLIND_DOCX))
    blob = "\n".join(p.text or "" for p in Document(str(BLIND_DOCX)).paragraphs)
    for leak in ("Khanh-Trinh", "utehy", "ioit", "github.com", "trinhnkt", "Hung Yen", "K.-T.N."):
        if leak.lower() in blob.lower():
            raise SystemExit(f"leak remains {leak}")
    for need in ("Anonymous Authors", "Author 1", "Omitted for double-blind review", "Grok 4.7", "GDPR"):
        if need not in blob:
            raise SystemExit(f"missing {need}")
    print("blind docx rebuilt", flush=True)
    ojs = HERE / "output" / "OJS_UPLOAD"
    shutil.copy2(BLIND_DOCX, REVIEW_DOCX)
    try:
        export_one(REVIEW_DOCX, BLIND_DOC, BLIND_PDF, "")
        shutil.copy2(REVIEW_DOCX, BLIND_DOCX)
    finally:
        for leftover in (
            REVIEW_DOCX,
            REVIEW_DOCX.with_suffix(".doc"),
            REVIEW_DOCX.with_suffix(".pdf"),
        ):
            if leftover.exists():
                leftover.unlink()
        for name in ("_layout_review.docx", "_layout_review.doc", "_layout_review.pdf"):
            p = ojs / name
            if p.exists():
                p.unlink()


if __name__ == "__main__":
    main()
